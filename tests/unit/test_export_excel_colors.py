"""Tests for the per-employee color code on GET /api/export_excel.

Covers ONLY the single-schedule export endpoint (the "Exportar" button):
resolved employee colors (custom override or default-by-first-name) must be
applied consistently to the schedule name/shift cells, the FORMATO column
(col K), and the "OBLIGACIONES / LIMPIEZA" name cells — while the fixed
LIBRE/VACACIONES/PERMISO status-cell styles must stay untouched.
"""
import datetime
import json
import os
import sys

import openpyxl
import pytest

_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
_PLAN = os.path.join(_ROOT, "planillas")
_BACKEND = os.path.join(_ROOT, "backend")
for _p in (_PLAN, _BACKEND):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import database as db  # noqa: E402

JEISON = "Jeison Aleman Tijerino"  # matches the "Jeison" default by leading-name
CUSTOM = "Custom Persona"  # has a saved custom override

SCHEDULE = {
    JEISON: {
        "Vie": "D_6-14", "Sáb": "OFF", "Dom": "VAC", "Lun": "PERM",
        "Mar": "D_6-14", "Mié": "D_6-14", "Jue": "D_6-14",
    },
    CUSTOM: {
        "Vie": "D_6-14", "Sáb": "D_6-14", "Dom": "D_6-14", "Lun": "D_6-14",
        "Mar": "D_6-14", "Mié": "D_6-14", "Jue": "D_6-14",
    },
}
TASKS = {
    JEISON: {"Vie": "Baños ↓PM"},
}
CUSTOM_BG = "AA11BB"
CUSTOM_FONT = "00FF00"


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DB_FILE", str(tmp_path / "test.db"))
    db.init_db()
    conn = db.get_conn()
    conn.execute(
        "INSERT INTO horarios_generados (nombre, horario, tareas, metadata, timestamp) "
        "VALUES (?, ?, ?, ?, ?)",
        (
            "Semana Test Colores",
            json.dumps(SCHEDULE, ensure_ascii=False),
            json.dumps(TASKS, ensure_ascii=False),
            json.dumps({}, ensure_ascii=False),
            datetime.datetime.now().isoformat(),
        ),
    )
    conn.execute("INSERT OR IGNORE INTO horario_config (id) VALUES (1)")
    conn.execute(
        "UPDATE horario_config SET excel_colors_json=? WHERE id=1",
        (json.dumps({CUSTOM: {"bg": CUSTOM_BG, "font": CUSTOM_FONT}}, ensure_ascii=False),),
    )
    conn.commit()
    conn.close()

    from fastapi.testclient import TestClient
    import main as app_main  # noqa: E402

    return TestClient(app_main.app)


def _rgb(color) -> str:
    """openpyxl round-trips colors as 8-char ARGB (e.g. '00663300'); keep the last 6."""
    raw = getattr(color, "rgb", None) or ""
    return str(raw)[-6:].upper()


def _load_sheet(response_content, tmp_path):
    out_path = tmp_path / "export_test.xlsm"
    out_path.write_bytes(response_content)
    wb = openpyxl.load_workbook(str(out_path))
    return wb["Horario"]


def _find_row_by_col1_value(ws, value):
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == value:
            return row
    raise AssertionError(f"row with value {value!r} not found")


class TestScheduleCells:
    def test_default_employee_name_cell_uses_resolved_default_color(self, client, tmp_path):
        res = client.get("/api/export_excel")
        assert res.status_code == 200
        ws = _load_sheet(res.content, tmp_path)

        row = _find_row_by_col1_value(ws, JEISON)
        name_cell = ws.cell(row=row, column=1)
        assert _rgb(name_cell.fill.fgColor) == "663300"
        assert _rgb(name_cell.font.color) == "FFFFFF"

    def test_custom_employee_name_cell_uses_custom_color(self, client, tmp_path):
        res = client.get("/api/export_excel")
        ws = _load_sheet(res.content, tmp_path)

        row = _find_row_by_col1_value(ws, CUSTOM)
        name_cell = ws.cell(row=row, column=1)
        assert _rgb(name_cell.fill.fgColor) == CUSTOM_BG
        assert _rgb(name_cell.font.color) == CUSTOM_FONT

    def test_off_cell_keeps_libre_style_regardless_of_employee_color(self, client, tmp_path):
        res = client.get("/api/export_excel")
        ws = _load_sheet(res.content, tmp_path)

        row = _find_row_by_col1_value(ws, JEISON)
        off_cell = ws.cell(row=row, column=3)  # Sáb -> OFF
        assert _rgb(off_cell.fill.fgColor) == "FFFF00"
        assert _rgb(off_cell.font.color) == "999999"

    def test_vac_and_perm_cells_keep_their_fixed_style(self, client, tmp_path):
        res = client.get("/api/export_excel")
        ws = _load_sheet(res.content, tmp_path)

        row = _find_row_by_col1_value(ws, JEISON)
        vac_cell = ws.cell(row=row, column=4)  # Dom -> VAC
        perm_cell = ws.cell(row=row, column=5)  # Lun -> PERM
        assert _rgb(vac_cell.fill.fgColor) == "C6EFCE"
        assert _rgb(vac_cell.font.color) == "006100"
        assert _rgb(perm_cell.fill.fgColor) == "FCE4D6"
        assert _rgb(perm_cell.font.color) == "9A3412"


class TestFormatoColumn:
    def test_formato_column_matches_resolved_colors(self, client, tmp_path):
        res = client.get("/api/export_excel")
        ws = _load_sheet(res.content, tmp_path)

        formato_col = 11  # column K
        jeison_row = _find_row_by_col1_value(ws, JEISON)
        fmt_cell = ws.cell(row=jeison_row, column=formato_col)
        assert fmt_cell.value == JEISON
        assert _rgb(fmt_cell.fill.fgColor) == "663300"
        assert _rgb(fmt_cell.font.color) == "FFFFFF"

        custom_row = _find_row_by_col1_value(ws, CUSTOM)
        fmt_cell_custom = ws.cell(row=custom_row, column=formato_col)
        assert fmt_cell_custom.value == CUSTOM
        assert _rgb(fmt_cell_custom.fill.fgColor) == CUSTOM_BG
        assert _rgb(fmt_cell_custom.font.color) == CUSTOM_FONT


class TestObligacionesSection:
    def test_obligaciones_name_cell_matches_resolved_colors(self, client, tmp_path):
        res = client.get("/api/export_excel")
        ws = _load_sheet(res.content, tmp_path)

        header_row = _find_row_by_col1_value(ws, "OBLIGACIONES / LIMPIEZA")
        task_header_row = header_row + 1
        jeison_task_row = task_header_row + 1  # first employee in dict order

        assert ws.cell(row=jeison_task_row, column=1).value == JEISON
        name_cell = ws.cell(row=jeison_task_row, column=1)
        assert _rgb(name_cell.fill.fgColor) == "663300"
        assert _rgb(name_cell.font.color) == "FFFFFF"

        custom_task_row = task_header_row + 2
        assert ws.cell(row=custom_task_row, column=1).value == CUSTOM
        custom_name_cell = ws.cell(row=custom_task_row, column=1)
        assert _rgb(custom_name_cell.fill.fgColor) == CUSTOM_BG
        assert _rgb(custom_name_cell.font.color) == CUSTOM_FONT
