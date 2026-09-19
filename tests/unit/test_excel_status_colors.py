"""Tests for configurable LIBRE / VACACIONES / PERMISO colors in the Excel export.

Status colors live in horario_config.excel_status_colors_json (isolated like
excel_colors_json) and apply ONLY to GET /api/export_excel.
"""
import datetime
import json
import os
import sys

import openpyxl
import pytest

_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
_PLAN = os.path.join(_ROOT, "planillas")
_BACKEND = os.path.join(_ROOT, "backend")
for _p in (_PLAN, _BACKEND):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import database as db  # noqa: E402
from excel_colors import resolve_excel_status_colors  # noqa: E402

EMP = "Jeison Aleman Tijerino"
SCHEDULE = {
    EMP: {
        "Vie": "D_6-14", "Sáb": "OFF", "Dom": "VAC", "Lun": "PERM",
        "Mar": "D_6-14", "Mié": "D_6-14", "Jue": "D_6-14",
    },
}


# ── Resolver ────────────────────────────────────────────────────────────────

class TestResolveStatusColors:
    def test_defaults_match_master_workbook(self):
        result = resolve_excel_status_colors({})
        assert list(result.keys()) == ["OFF", "VAC", "PERM"]
        assert result["OFF"] == {"label": "LIBRE", "bg": "FFFF00", "font": "999999", "source": "default"}
        assert result["VAC"] == {"label": "VACACIONES", "bg": "C6EFCE", "font": "006100", "source": "default"}
        assert result["PERM"] == {"label": "PERMISO", "bg": "FCE4D6", "font": "9A3412", "source": "default"}

    def test_custom_overrides_default(self):
        result = resolve_excel_status_colors({"VAC": {"bg": "#123456", "font": "ffffff"}})
        assert result["VAC"]["bg"] == "123456"
        assert result["VAC"]["font"] == "FFFFFF"
        assert result["VAC"]["source"] == "custom"
        assert result["OFF"]["source"] == "default"

    def test_custom_bg_only_gets_auto_contrast_font(self):
        result = resolve_excel_status_colors({"OFF": {"bg": "000000"}})
        assert result["OFF"]["bg"] == "000000"
        assert result["OFF"]["font"] == "FFFFFF"

    def test_invalid_hex_and_unknown_codes_are_ignored(self):
        result = resolve_excel_status_colors({"PERM": {"bg": "nope"}, "XYZ": {"bg": "112233"}})
        assert result["PERM"]["source"] == "default"
        assert "XYZ" not in result


# ── API + export ────────────────────────────────────────────────────────────

@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DB_FILE", str(tmp_path / "test.db"))
    db.init_db()
    conn = db.get_conn()
    conn.execute(
        "INSERT INTO empleados (nombre, tipo_pago, activo) VALUES (?, 'tarjeta', 1)",
        (EMP,),
    )
    conn.execute("INSERT INTO horario_empleados (nombre) VALUES (?)", (EMP,))
    conn.execute(
        "INSERT INTO horarios_generados (nombre, horario, tareas, metadata, timestamp) "
        "VALUES (?, ?, ?, ?, ?)",
        (
            "Semana Test Estados",
            json.dumps(SCHEDULE, ensure_ascii=False),
            json.dumps({}, ensure_ascii=False),
            json.dumps({}, ensure_ascii=False),
            datetime.datetime.now().isoformat(),
        ),
    )
    conn.commit()
    conn.close()

    from fastapi.testclient import TestClient
    import main as app_main  # noqa: E402

    return TestClient(app_main.app)


def _status(payload, code):
    return next(s for s in payload["statuses"] if s["code"] == code)


def _rgb(color) -> str:
    raw = getattr(color, "rgb", None) or ""
    return str(raw)[-6:].upper()


def _load_sheet(content, tmp_path):
    out_path = tmp_path / "export_status.xlsm"
    out_path.write_bytes(content)
    return openpyxl.load_workbook(str(out_path))["Horario"]


def _find_row(ws, value, column=1):
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=column).value == value:
            return row
    raise AssertionError(f"{value!r} not found")


def test_get_returns_statuses_with_defaults(client):
    data = client.get("/api/excel-colors").json()
    assert [s["code"] for s in data["statuses"]] == ["OFF", "VAC", "PERM"]
    off = _status(data, "OFF")
    assert off["label"] == "LIBRE"
    assert off["bg"] == "FFFF00"
    assert off["source"] == "default"


def test_put_status_colors_roundtrip(client):
    res = client.put(
        "/api/excel-colors",
        json={"colors": {}, "status_colors": {"VAC": {"bg": "0000FF", "font": "FFFFFF"}}},
    )
    assert res.status_code == 200
    vac = _status(res.json(), "VAC")
    assert (vac["bg"], vac["font"], vac["source"]) == ("0000FF", "FFFFFF", "custom")
    assert _status(client.get("/api/excel-colors").json(), "VAC")["source"] == "custom"


def test_put_without_status_colors_keeps_stored_ones(client):
    client.put("/api/excel-colors", json={"colors": {}, "status_colors": {"OFF": {"bg": "EEEEEE"}}})
    client.put("/api/excel-colors", json={"colors": {EMP: {"bg": "112233"}}})
    assert _status(client.get("/api/excel-colors").json(), "OFF")["bg"] == "EEEEEE"


def test_put_empty_status_colors_resets_to_default(client):
    client.put("/api/excel-colors", json={"colors": {}, "status_colors": {"OFF": {"bg": "EEEEEE"}}})
    res = client.put("/api/excel-colors", json={"colors": {}, "status_colors": {}})
    assert _status(res.json(), "OFF")["bg"] == "FFFF00"


def test_post_config_does_not_wipe_status_colors(client):
    client.put("/api/excel-colors", json={"colors": {}, "status_colors": {"PERM": {"bg": "ABCDEF"}}})
    cfg = client.get("/api/config").json()
    assert client.post("/api/config", json=cfg).status_code == 200
    assert _status(client.get("/api/excel-colors").json(), "PERM")["bg"] == "ABCDEF"


def test_export_applies_custom_status_colors(client, tmp_path):
    client.put(
        "/api/excel-colors",
        json={
            "colors": {},
            "status_colors": {
                "OFF": {"bg": "111111", "font": "EEEEEE"},
                "VAC": {"bg": "222222", "font": "DDDDDD"},
                "PERM": {"bg": "333333", "font": "CCCCCC"},
            },
        },
    )
    res = client.get("/api/export_excel")
    assert res.status_code == 200
    ws = _load_sheet(res.content, tmp_path)

    row = _find_row(ws, EMP)
    off_cell, vac_cell, perm_cell = (ws.cell(row=row, column=c) for c in (3, 4, 5))
    assert (_rgb(off_cell.fill.fgColor), _rgb(off_cell.font.color)) == ("111111", "EEEEEE")
    assert (_rgb(vac_cell.fill.fgColor), _rgb(vac_cell.font.color)) == ("222222", "DDDDDD")
    assert (_rgb(perm_cell.fill.fgColor), _rgb(perm_cell.font.color)) == ("333333", "CCCCCC")

    # FORMATO legend (column K) LIBRE cell follows the same custom color
    libre_row = _find_row(ws, "LIBRE", column=11)
    libre_cell = ws.cell(row=libre_row, column=11)
    assert (_rgb(libre_cell.fill.fgColor), _rgb(libre_cell.font.color)) == ("111111", "EEEEEE")
